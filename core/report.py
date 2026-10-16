"""Exportación de resultados a PDF, LaTeX y Excel sin acceder a la red."""

from fractions import Fraction
from html import escape
from pathlib import Path

from core.models import PasoSimplex, Problema, ResultadoSimplex
from core.numeros import NumM
from utils.paleta import COLOR_COLUMNA, COLOR_FILA, COLOR_PIVOTE, COLOR_Z


def _filas(paso: PasoSimplex) -> list[list[str]]:
    filas = [["Cj", ""] + [str(c) for c in paso.cj] + [""]]
    filas += [
        [n, str(c)] + [str(v) for v in fila]
        for n, c, fila in zip(paso.base, paso.cb, paso.tabla)
    ]
    filas += [["Zj − Cj", ""] + [str(z) for z in paso.fila_z]]
    return filas


def _color(paso: PasoSimplex, i: int, j: int) -> tuple[str, bool]:
    fila = paso.fila_pivote is not None and i == paso.fila_pivote + 1
    columna = paso.col_pivote is not None and j == paso.col_pivote + 2
    return (
        (
            COLOR_PIVOTE
            if fila and columna
            else (
                COLOR_FILA
                if fila
                else (
                    COLOR_COLUMNA
                    if columna
                    else COLOR_Z if i == len(paso.tabla) + 1 else "#FFFFFF"
                )
            )
        ),
        fila and columna,
    )


def _expresion(coefs: list[Fraction], nombres: list[str]) -> str:
    return " ".join(
        f"{'+' if c >= 0 else '-'} {abs(c)}{n}" for c, n in zip(coefs, nombres)
    ).lstrip("+ ")


def _enunciado(p: Problema) -> list[str]:
    return (
        [f"{p.sentido} Z = {_expresion(p.coef_obj,p.nombres)}", "s.a."]
        + [f"{_expresion(r.coefs,p.nombres)} {r.tipo} {r.rhs}" for r in p.restricciones]
        + ["Todas las variables son ≥ 0."]
    )


def construir_html(
    problema: Problema,
    resultado: ResultadoSimplex,
    imagen_grafico: Path | None = None,
    metodo: str = "Simplex Gran M",
) -> str:
    """Compone un informe con problema, tablas marcadas, operaciones y solución."""
    partes = [
        "<html><head><meta charset='utf-8'></head><body>",
        "<h1>SimplexLab</h1>",
        f"<p>Método aplicado: {escape(metodo)}</p>",
        "<pre>" + escape("\n".join(_enunciado(problema))) + "</pre>",
    ]
    for paso in resultado.pasos:
        partes += [
            f"<h2 style='page-break-before: always'>Tabla {paso.numero}</h2>",
            f"<p>{escape(paso.mensaje)}</p>",
            "<table border='1' cellspacing='0' cellpadding='5'><tr>",
        ]
        partes.extend(
            f"<th>{escape(n)}</th>" for n in ["Base", "Cb"] + paso.encabezados
        )
        partes.append("</tr>")
        for i, fila in enumerate(_filas(paso)):
            partes.append("<tr>")
            for j, valor in enumerate(fila):
                color, negrita = _color(paso, i, j)
                texto = escape(valor)
                if negrita:
                    texto = f"<b>{texto}</b>"
                partes.append(f"<td bgcolor='{color}'>{texto}</td>")
            partes.append("</tr>")
        partes += ["</table><h3>Operaciones entre filas</h3><ul>"]
        partes.extend(f"<li>{escape(o)}</li>" for o in paso.operaciones)
        partes.append("</ul>")
    partes += [
        "<h2>Resultado final</h2>",
        f"<p>Estado: {escape(resultado.estado)}; Z* = {resultado.z}</p>",
    ]
    if resultado.solucion:
        partes.append(
            "<p>"
            + "; ".join(escape(f"{n} = {v}") for n, v in resultado.solucion.items())
            + "</p>"
        )
    partes.append(
        f"<p>Iteraciones: {len(resultado.pasos)-1}. Degenerado: "
        f"{'sí' if resultado.degenerado else 'no'}.</p>"
    )
    partes.extend(f"<p>{escape(a)}</p>" for a in resultado.advertencias)
    if imagen_grafico:
        partes.append(
            f"<h2 style='page-break-before: always'>Método "
            f"gráfico</h2><img width='650' "
            f"src='{escape(imagen_grafico.resolve().as_uri())}'>"
        )
    return "".join(partes) + "</body></html>"


def exportar_pdf(
    problema: Problema,
    resultado: ResultadoSimplex,
    ruta: Path,
    imagen_grafico: Path | None = None,
    metodo: str = "Simplex Gran M",
) -> None:
    """Imprime el informe mediante QTextDocument y QPrinter.

    Raises:
        ValueError: Si no existe una aplicación Qt para imprimir el documento.
    """
    from PySide6.QtGui import QGuiApplication, QPageSize, QTextDocument
    from PySide6.QtPrintSupport import QPrinter, QPrinterInfo

    if QGuiApplication.instance() is None:
        raise ValueError("La exportación PDF requiere una aplicación Qt activa.")
    impresora = QPrinter(QPrinterInfo(), QPrinter.PrinterMode.HighResolution)
    impresora.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
    impresora.setOutputFileName(str(ruta))
    impresora.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    documento = QTextDocument()
    documento.setHtml(construir_html(problema, resultado, imagen_grafico, metodo))
    documento.print_(impresora)
    if not ruta.is_file() or ruta.stat().st_size == 0:
        raise OSError("No se pudo generar el PDF en la ruta seleccionada.")


def _latex_texto(texto: str) -> str:
    mapa = {
        "\\": r"\textbackslash{}",
        "_": r"\_",
        "%": r"\%",
        "&": r"\&",
        "#": r"\#",
        "$": r"\$",
        "{": r"\{",
        "}": r"\}",
        "−": "-",
        "←": r"$\leftarrow$",
        "·": r"$\cdot$",
        "≥": r"$\geq$",
        "≤": r"$\leq$",
        "⚠": "Advertencia",
    }
    return "".join(mapa.get(c, c) for c in texto)


def _latex_numero(valor: Fraction | NumM) -> str:
    if isinstance(valor, NumM):
        if not valor.m:
            return _latex_numero(valor.c)
        termino = "M" if abs(valor.m) == 1 else _latex_numero(abs(valor.m)) + "M"
        termino = ("-" if valor.m < 0 else "") + termino
        return termino + (
            (" + " if valor.c > 0 else " - ") + _latex_numero(abs(valor.c))
            if valor.c
            else ""
        )
    return (
        str(valor.numerator)
        if valor.denominator == 1
        else rf"\frac{{{valor.numerator}}}{{{valor.denominator}}}"
    )


def exportar_latex(problema: Problema, resultado: ResultadoSimplex, ruta: Path) -> None:
    """Genera un documento compilable con amsmath, fracciones y tablas."""
    partes = [
        r"\documentclass{article}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage{amsmath}",
        r"\usepackage[margin=1.5cm]{geometry}",
        r"\begin{document}",
        r"\section*{SimplexLab}",
    ]
    partes.extend(_latex_texto(s) + r"\par" for s in _enunciado(problema))
    for paso in resultado.pasos:
        partes.extend(
            [
                rf"\section*{{Tabla {paso.numero}}}",
                _latex_texto(paso.mensaje) + r"\par",
                r"\begin{tabular}{" + "l" * (len(paso.encabezados) + 2) + "}",
                " & ".join(_latex_texto(n) for n in ["Base", "Cb"] + paso.encabezados)
                + r" \\ \hline",
            ]
        )
        filas = [["Cj", ""] + paso.cj + [""]]
        filas += [[n, c] + f for n, c, f in zip(paso.base, paso.cb, paso.tabla)]
        filas += [["Zj-Cj", ""] + paso.fila_z]
        for fila in filas:
            partes.append(
                " & ".join(
                    (
                        "$" + _latex_numero(v) + "$"
                        if isinstance(v, (Fraction, NumM))
                        else _latex_texto(v)
                    )
                    for v in fila
                )
                + r" \\"
            )
        partes.append(r"\end{tabular}\par")
        partes.extend(_latex_texto(o) + r"\par" for o in paso.operaciones)
    partes.append(r"\section*{Resultado final}")
    partes.append(_latex_texto(f"Estado: {resultado.estado}") + r"\par")
    if resultado.z is not None:
        partes.append("$Z^* = " + _latex_numero(resultado.z) + "$" + r"\par")
        partes.extend(
            _latex_texto(n) + " $= " + _latex_numero(v) + "$" + r"\par"
            for n, v in resultado.solucion.items()
        )
    partes.extend(_latex_texto(a) + r"\par" for a in resultado.advertencias)
    partes.append(r"\end{document}")
    ruta.write_text("\n".join(partes), encoding="utf-8")


def exportar_excel(problema: Problema, resultado: ResultadoSimplex, ruta: Path) -> None:
    """Guarda una hoja por tabla con colores y fracciones como texto exacto."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    resumen = libro.active
    resumen.title = "Problema y resultado"
    for linea in _enunciado(problema):
        resumen.append([linea])
    resumen.append(["Estado", resultado.estado])
    resumen.append(["Z*", str(resultado.z) if resultado.z is not None else ""])
    resumen.append(["Iteraciones", len(resultado.pasos) - 1])
    resumen.append(["Degenerado", str(resultado.degenerado)])
    for n, v in (resultado.solucion or {}).items():
        resumen.append([n, str(v)])
    for a in resultado.advertencias:
        resumen.append([a])
    resumen.column_dimensions["A"].width = 70
    for paso in resultado.pasos:
        hoja = libro.create_sheet(f"Paso {paso.numero}")
        hoja.append(["Base", "Cb"] + paso.encabezados)
        for i, fila in enumerate(_filas(paso)):
            hoja.append(fila)
            for j in range(len(fila)):
                color, negrita = _color(paso, i, j)
                celda = hoja.cell(i + 2, j + 1)
                # Las fracciones y nombres son datos, nunca fórmulas de Excel.
                celda.data_type = "s"
                celda.fill = PatternFill("solid", fgColor=color.lstrip("#"))
                celda.font = Font(bold=negrita)
        hoja.append([])
        hoja.append([paso.mensaje])
        for operacion in paso.operaciones:
            hoja.append([operacion])
        for j in range(1, len(paso.encabezados) + 3):
            hoja.column_dimensions[get_column_letter(j)].width = 18
        hoja.freeze_panes = "C3"
    for hoja in libro:
        for fila in hoja:
            for celda in fila:
                if isinstance(celda.value, str):
                    celda.data_type = "s"
    libro.save(ruta)
