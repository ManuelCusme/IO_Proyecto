"""Exportadores de archivos temporales, sin credenciales ni acceso a red."""

from pathlib import Path

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from core.report import construir_html, exportar_excel, exportar_latex, exportar_pdf
from core.simplex_solver import resolver_simplex
from tests.casos_taha import CASOS


def test_exportadores(qapp: QApplication, tmp_path: Path) -> None:
    """Genera los tres formatos y reabre Excel para verificar valores y pivote."""
    p = CASOS[0][0]
    r = resolver_simplex(p)
    for nombre, exportador in [
        ("pdf", exportar_pdf),
        ("tex", exportar_latex),
        ("xlsx", exportar_excel),
    ]:
        ruta = tmp_path / f"resultado.{nombre}"
        exportador(p, r, ruta)
        assert ruta.stat().st_size > 100
    assert (tmp_path / "resultado.pdf").read_bytes().startswith(b"%PDF")
    tex = (tmp_path / "resultado.tex").read_text(encoding="utf-8")
    assert r"\frac{1}{2}" in tex and r"\begin{tabular}" in tex
    libro = load_workbook(tmp_path / "resultado.xlsx")
    assert libro.sheetnames == ["Problema y resultado", "Paso 0", "Paso 1", "Paso 2"]
    pivote = libro["Paso 0"].cell(4, 4)
    assert pivote.value == "2" and pivote.font.bold
    assert pivote.fill.fgColor.rgb.endswith("FFE082")
    assert "#FFE082" in construir_html(p, r)
    assert "FZ" in construir_html(p, r)
    libro.close()


def test_pdf_con_grafico(qapp: QApplication, tmp_path: Path) -> None:
    """La exportación desde la ventana incorpora la imagen del canvas."""
    from gui.main_window import MainWindow

    ventana = MainWindow()
    ventana.panel.cargar_problema(CASOS[0][0])
    ventana.panel.resolver_boton.click()
    destino = tmp_path / "con-grafico.pdf"
    ventana.exportar_a(destino, "pdf")
    assert destino.read_bytes().startswith(b"%PDF")
    assert destino.stat().st_size > 5000
    ventana.close()
